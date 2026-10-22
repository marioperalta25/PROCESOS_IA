"""Lector de flujos por bloques para archivos Excel (XLSX) con bajo consumo de RAM."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl

from procesos_ia.dominio.excepciones import ErrorProcesamiento
from procesos_ia.dominio.modelos import RegistroEntrada
from procesos_ia.infraestructura.archivos.detector_formato import detectar_encabezados


class LectorXlsx:
    """Implementa LectorFlujoProtocol para libros Excel en modo solo lectura por flujo."""

    def __init__(self, nombre_hoja: str | None = None) -> None:
        self._nombre_hoja = nombre_hoja
        self._ruta: Path | None = None
        self._encabezados: list[str] = []
        self._filas_saltar_cabecera: int = 0
        self._total_estimado: int | None = None
        self._libro_flujo: Any | None = None
        self._iter_flujo: Any | None = None
        self._contador_flujo: int = 0

    @property
    def encabezados(self) -> list[str]:
        return list(self._encabezados)

    def inicializar(self, ruta_archivo: Path) -> list[str]:
        """Abre el libro en modo flujo, inspecciona las primeras 5 filas y extrae cabecera."""
        if not ruta_archivo.is_file():
            raise ErrorProcesamiento(f"El archivo no existe: {ruta_archivo}")

        self._ruta = ruta_archivo
        libro = openpyxl.load_workbook(
            filename=str(ruta_archivo), read_only=True, data_only=True
        )
        try:
            hoja = (
                libro[self._nombre_hoja]
                if self._nombre_hoja and self._nombre_hoja in libro.sheetnames
                else libro.active
            )
            filas_muestra: list[list[str]] = []
            for fila in hoja.iter_rows(values_only=True):
                fila_texto = [
                    "" if c is None else str(c).strip() for c in fila
                ]
                filas_muestra.append(fila_texto)
                if len(filas_muestra) >= 5:
                    break

            inspeccion = detectar_encabezados(filas_muestra)
            if inspeccion.tiene_encabezado:
                self._encabezados = inspeccion.encabezados
                self._filas_saltar_cabecera = inspeccion.indice_fila_encabezado + 1
            else:
                self._encabezados = inspeccion.encabezados
                self._filas_saltar_cabecera = inspeccion.indice_fila_encabezado

            return list(self._encabezados)
        finally:
            libro.close()

    def abrir_flujo(self, saltar_hasta_indice: int) -> None:
        """Abre el libro una sola vez y avanza el cursor al punto de reanudacion."""
        if self._ruta is None:
            raise ErrorProcesamiento("El lector no ha sido inicializado")

        self.cerrar()
        self._libro_flujo = openpyxl.load_workbook(
            filename=str(self._ruta), read_only=True, data_only=True
        )
        hoja = (
            self._libro_flujo[self._nombre_hoja]
            if self._nombre_hoja and self._nombre_hoja in self._libro_flujo.sheetnames
            else self._libro_flujo.active
        )
        self._iter_flujo = hoja.iter_rows(values_only=True)

        for _ in range(self._filas_saltar_cabecera):
            try:
                next(self._iter_flujo)
            except StopIteration:
                return

        self._contador_flujo = 0
        while self._contador_flujo < saltar_hasta_indice:
            try:
                fila = next(self._iter_flujo)
                fila_texto = ["" if c is None else str(c).strip() for c in fila]
                if any(fila_texto):
                    self._contador_flujo += 1
            except StopIteration:
                break

    def leer_siguiente_bloque(self, tamano_bloque: int) -> list[RegistroEntrada]:
        """Extrae el siguiente bloque desde el flujo Excel abierto en tiempo O(1) por fila."""
        if self._iter_flujo is None or self._ruta is None:
            raise ErrorProcesamiento("El flujo no esta abierto; llame a abrir_flujo()")

        nombre_archivo = self._ruta.name
        registros: list[RegistroEntrada] = []

        while len(registros) < tamano_bloque:
            try:
                fila = next(self._iter_flujo)
            except StopIteration:
                break

            fila_texto = ["" if c is None else str(c).strip() for c in fila]
            if not any(fila_texto):
                continue

            self._contador_flujo += 1
            valores_originales: dict[str, str] = {}
            for idx, col_nombre in enumerate(self._encabezados):
                val = fila_texto[idx] if idx < len(fila_texto) else ""
                valores_originales[col_nombre] = val

            identificador = f"{nombre_archivo}:{self._contador_flujo}"
            registros.append(
                RegistroEntrada(
                    identificador=identificador,
                    indice_fila=self._contador_flujo,
                    valores_originales=valores_originales,
                    datos_para_modelo=dict(valores_originales),
                )
            )

        return registros

    def leer_bloque(
        self,
        tamano_bloque: int,
        saltar_hasta_indice: int,
    ) -> Iterator[RegistroEntrada]:
        """Itera filas por flujo sin cargar el libro completo en memoria."""
        if self._ruta is None:
            raise ErrorProcesamiento("El lector no ha sido inicializado")

        nombre_archivo = self._ruta.name
        libro = openpyxl.load_workbook(
            filename=str(self._ruta), read_only=True, data_only=True
        )
        try:
            hoja = (
                libro[self._nombre_hoja]
                if self._nombre_hoja and self._nombre_hoja in libro.sheetnames
                else libro.active
            )
            contador_registro = 0
            emitidos = 0
            filas_iter = hoja.iter_rows(values_only=True)

            for _ in range(self._filas_saltar_cabecera):
                try:
                    next(filas_iter)
                except StopIteration:
                    return

            for fila in filas_iter:
                fila_texto = [
                    "" if c is None else str(c).strip() for c in fila
                ]
                if not any(fila_texto):
                    continue

                contador_registro += 1
                if contador_registro <= saltar_hasta_indice:
                    continue

                valores_originales: dict[str, str] = {}
                for idx, col_nombre in enumerate(self._encabezados):
                    val = fila_texto[idx] if idx < len(fila_texto) else ""
                    valores_originales[col_nombre] = val

                identificador = f"{nombre_archivo}:{contador_registro}"
                yield RegistroEntrada(
                    identificador=identificador,
                    indice_fila=contador_registro,
                    valores_originales=valores_originales,
                    datos_para_modelo=dict(valores_originales),
                )

                emitidos += 1
                if emitidos >= tamano_bloque:
                    break
        finally:
            libro.close()

    def estimar_total_registros(self) -> int:
        """Estima o cuenta filas totales mediante el descriptor de hoja."""
        if self._ruta is None:
            return 0
        if self._total_estimado is not None:
            return self._total_estimado

        libro = openpyxl.load_workbook(
            filename=str(self._ruta), read_only=True, data_only=True
        )
        try:
            hoja = (
                libro[self._nombre_hoja]
                if self._nombre_hoja and self._nombre_hoja in libro.sheetnames
                else libro.active
            )
            if hoja.max_row is not None and hoja.max_row > self._filas_saltar_cabecera:
                self._total_estimado = hoja.max_row - self._filas_saltar_cabecera
                return self._total_estimado

            total = 0
            filas_iter = hoja.iter_rows(values_only=True)
            for _ in range(self._filas_saltar_cabecera):
                try:
                    next(filas_iter)
                except StopIteration:
                    break
            for fila in filas_iter:
                if any(c is not None and str(c).strip() for c in fila):
                    total += 1
            self._total_estimado = total
            return total
        finally:
            libro.close()

    def cerrar(self) -> None:
        """Libera referencias locales y cierra el libro de flujo si esta abierto."""
        if self._libro_flujo is not None:
            self._libro_flujo.close()
            self._libro_flujo = None
        self._iter_flujo = None
        self._total_estimado = None
